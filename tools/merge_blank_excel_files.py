from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


INPUT_FILES = [
    Path("Данные для работы") / "Лист Microsoft Excel.xlsx",
    Path("Данные для работы") / "Сопоставление_заказа_заготовок (1).xlsx",
    Path("Данные для работы") / "Исходные данные для изготовления_заказа заготовок.xlsx",
]
OUTPUT_DIR = Path("reports") / "blank_excel_merge_20260728"
OUTPUT_FILE = OUTPUT_DIR / "Итоговая_библиотека_заготовок.xlsx"

TARGET_COLUMNS = [
    "Источник файла",
    "IPS",
    "Обозначение",
    "Наименование",
    "Вид заготовки",
    "Заготовка",
    "Размер заготовки",
    "Материал",
    "Гост материала",
    "Гост сортамента",
    "Количество (норма расхода материала)",
    "Ед. измерения",
    "УТ код",
    "Номенклатура НСИ",
    "Комментарий",
]

LONG_TYPES = {"Круг", "Труба", "Квадрат", "Шестигранник", "Пруток"}


def text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    value = str(value).strip()
    return re.sub(r"\s+", " ", value)


def decimal_text(value: Any) -> str:
    if value is None or pd.isna(value) or value == "":
        return "0"
    if isinstance(value, str):
        match = re.search(r"-?\d+(?:[,.]\d+)?", value)
        if not match:
            return "0"
        number = float(match.group(0).replace(",", "."))
    else:
        number = float(value)
    return f"{number:.6f}".rstrip("0").rstrip(".").replace(".", ",")


def norm_gost(value: str) -> str:
    match = re.search(r"ГОСТ\s+[РA-ZА-Яа-я0-9.\-\/ ]*?\d{2,5}(?:-\d{2,4})?", value, flags=re.IGNORECASE)
    return text(match.group(0)) if match else ""


def remove_gost(value: str) -> str:
    return text(re.sub(r"ГОСТ\s+[РA-ZА-Яа-я0-9.\-\/ ]*?\d{2,5}(?:-\d{2,4})?", "", value, flags=re.IGNORECASE).replace(",", " "))


def split_material(value: Any) -> tuple[str, str]:
    source = text(value)
    return remove_gost(source), norm_gost(source)


def clean_profile(value: str) -> str:
    value = text(value)
    value = re.sub(r"(\d+(?:[,.]\d+)?)\s*[-–]\s*[А-ЯA-Z]\d?\b", r"\1", value, flags=re.IGNORECASE)
    value = re.sub(r"\b[А-ЯA-Z]\d?\b$", "", value, flags=re.IGNORECASE)
    return text(value)


def parse_blank_name(value: Any) -> tuple[str, str, str, str]:
    raw = text(value)
    sort_gost = norm_gost(raw)
    no_gost = remove_gost(raw)
    material_match = re.search(r"(Сталь\s+[0-9A-Za-zА-Яа-я]+|СЧ\d+|Бр[А-Яа-яA-Za-z0-9\-]+)", no_gost, flags=re.IGNORECASE)
    material = text(material_match.group(1)) if material_match else ""
    if material:
        no_gost = text(no_gost.replace(material_match.group(0), ""))

    kind = infer_kind(no_gost)
    size = parse_size(no_gost, kind)
    blank = clean_profile(remove_size_text(no_gost))
    if not blank and kind:
        blank = kind
    return kind, blank, size, material


def infer_kind(value: Any) -> str:
    raw = text(value).lower()
    if "труб" in raw:
        return "Труба"
    if "круг" in raw or "ø" in raw or re.search(r"\bф\s*\d+", raw):
        return "Круг"
    if "квадрат" in raw:
        return "Квадрат"
    if "лист" in raw:
        return "Лист"
    if "плит" in raw:
        return "Плита"
    if "поков" in raw:
        return "Поковка"
    if "отлив" in raw or "лить" in raw:
        return "Литье"
    return ""


def remove_size_text(value: str) -> str:
    value = re.sub(r"[ØФф]\s*\d+(?:[,.]\d+)?\s*(?:мм)?", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\bD\s*\d+(?:[,.]\d+)?\b", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\bL\s*=?\s*\d+(?:[,.]\d+)?\s*(?:мм)?", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\d+(?:[,.]\d+)?\s*[xх×]\s*\d+(?:[,.]\d+)?(?:\s*[xх×]\s*\d+(?:[,.]\d+)?)?\s*(?:мм)?", "", value, flags=re.IGNORECASE)
    return text(value)


def parse_size(value: Any, kind: str) -> str:
    raw = text(value).replace("×", "х").replace("X", "х").replace("x", "х")
    diameter = re.search(r"(?:[ØФф]\s*|D\s*=?\s*|ф\s*)(\d+(?:[,.]\d+)?)", raw, flags=re.IGNORECASE)
    wall = re.search(r"(?:S|стенка|толщина стенки)\s*=?\s*(\d+(?:[,.]\d+)?)", raw, flags=re.IGNORECASE)
    numbers = [decimal_text(item) for item in re.findall(r"\d+(?:[,.]\d+)?", raw)]

    if kind == "Труба" and diameter:
        result = f"D{decimal_text(diameter.group(1))}"
        if wall:
            result += f" S{decimal_text(wall.group(1))}"
        elif len(numbers) >= 2:
            result += f" S{numbers[1]}"
        return result
    if kind in LONG_TYPES and diameter:
        return f"D{decimal_text(diameter.group(1))}"
    if kind in {"Лист", "Плита", "Поковка", "Литье", "Квадрат"} and len(numbers) >= 3:
        length, width, height = numbers[0], numbers[1], numbers[2]
        return f"W{height} H{width} L{length}"
    if len(numbers) >= 3 and "х" in raw:
        length, width, height = numbers[0], numbers[1], numbers[2]
        return f"W{height} H{width} L{length}"
    if len(numbers) == 1 and kind in LONG_TYPES:
        return f"D{numbers[0]}"
    return ""


def consumption_from_length(value: Any) -> str:
    number = decimal_text(value)
    return decimal_text(float(number.replace(",", ".")) / 1000) if number != "0" else "0"


def split_designation_name(value: Any) -> tuple[str, str]:
    raw = text(value)
    match = re.match(r"(.+?)\s*\((.+)\)\s*$", raw)
    if match:
        return text(match.group(1)), text(match.group(2))
    return raw, ""


def standard_row(source_file: str, **values: Any) -> dict[str, str]:
    row = {column: "" for column in TARGET_COLUMNS}
    row["Источник файла"] = source_file
    for key, value in values.items():
        row[key] = text(value) if key not in {"Количество (норма расхода материала)"} else decimal_text(value)
    for numeric in ["Количество (норма расхода материала)"]:
        row[numeric] = row[numeric] if row[numeric] else "0"
    if row["Ед. измерения"].lower().replace(" ", "") in {"пог.м", "п.м", "м"}:
        # For long products length is the consumption norm, not part of the size key.
        row["Размер заготовки"] = text(re.sub(r"\s*L\d+(?:[,.]\d+)?\b", "", row["Размер заготовки"], flags=re.IGNORECASE))
    return row


def load_prepared_library(path: Path) -> list[dict[str, str]]:
    df = pd.read_excel(path, sheet_name=0, dtype=str).fillna("")
    rows = []
    for _, source in df.iterrows():
        # In this source file the "Материал" and "УТ код" headers are shifted for several rows.
        raw_blank = text(source.get("Заготовка"))
        ut_code = text(source.get("УТ код"))
        material_text = text(source.get("Материал"))
        if material_text.startswith("УТ") and not ut_code.startswith("УТ"):
            ut_code, material_text = material_text, ut_code
        kind, blank, size, parsed_material = parse_blank_name(raw_blank)
        material = parsed_material or material_text
        material, material_gost = split_material(material)
        rows.append(
            standard_row(
                path.name,
                IPS=source.get("IPS"),
                Обозначение=source.get("Обозначение"),
                Наименование=source.get("Наименование"),
                **{
                    "Вид заготовки": kind,
                    "Заготовка": blank,
                    "Размер заготовки": size,
                    "Материал": material,
                    "Гост материала": source.get("Гост материала") or material_gost,
                    "Гост сортамента": source.get("Гост сортамента"),
                    "Количество (норма расхода материала)": source.get("Количество"),
                    "Ед. измерения": source.get("Ед. изм."),
                    "УТ код": ut_code,
                    "Номенклатура НСИ": raw_blank,
                    "Комментарий": "подготовленный лист",
                },
            )
        )
    return rows


def load_matching_result(path: Path) -> list[dict[str, str]]:
    df = pd.read_excel(path, sheet_name="Итог", dtype=str).fillna("")
    rows = []
    for _, source in df.iterrows():
        material, material_gost = split_material(source.get("Материал"))
        kind = text(source.get("Вид заготовки"))
        size = parse_size(source.get("размер заготовки"), kind)
        blank = f"{kind} {size.replace('D', '').split()[0]}" if kind == "Круг" and size.startswith("D") else kind
        rows.append(
            standard_row(
                path.name,
                IPS=source.get("IPS"),
                Обозначение=source.get("Обозначение"),
                Наименование=source.get("Наименование"),
                **{
                    "Вид заготовки": kind,
                    "Заготовка": blank,
                    "Размер заготовки": size,
                    "Материал": material,
                    "Гост материала": material_gost,
                    "Гост сортамента": "",
                    "Количество (норма расхода материала)": source.get("потребность в материале"),
                    "Ед. измерения": source.get("Ед. изм."),
                    "УТ код": source.get("УТ код"),
                    "Номенклатура НСИ": source.get("наименование Заготовки по 1с"),
                    "Комментарий": "лист Итог",
                },
            )
        )
    return rows


def load_raw_manufacturing(path: Path) -> list[dict[str, str]]:
    rows = []
    sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
    for sheet_name, raw in sheets.items():
        header_idx = raw.index[raw.apply(lambda row: row.astype(str).str.contains("IPS", case=False, na=False).any(), axis=1)]
        if len(header_idx) == 0:
            continue
        idx = int(header_idx[0])
        df = raw.iloc[idx + 1 :].copy()
        df.columns = [text(value) for value in raw.iloc[idx].tolist()]
        df = df.loc[:, [column for column in df.columns if column]]
        df = df.fillna("")

        for _, source in df.iterrows():
            ips = text(source.get("IPS"))
            if not ips:
                continue
            material, material_gost = split_material(source.get("Марка материала, нормативный документ"))
            sort_gost = text(source.get("Сортамент, нормативный документ"))

            if "труб" in sheet_name.lower():
                kind = "Труба"
                diameter = source.get("Диаметр заготовки трубы (Ø, мм)")
                wall = source.get("Толщина стенки (L,мм)")
                size = f"D{decimal_text(diameter)} S{decimal_text(wall)}" if decimal_text(diameter) != "0" else ""
                blank = f"Труба {size}" if size else "Труба"
                qty = consumption_from_length(source.get("Длина заготовки (L, мм)"))
                unit = "пог. м"
            elif "направля" in sheet_name.lower():
                designation, name = split_designation_name(source.get("НАИМЕНОВАНИЕ"))
                kind, blank, size, parsed_material = parse_blank_name(source.get("НАИМЕНОВАНИЕ ЗАГОТОВКИ"))
                material = parsed_material or material
                qty = source.get("ШТ")
                unit = "шт"
                rows.append(
                    standard_row(
                        path.name,
                        IPS=ips,
                        Обозначение=designation,
                        Наименование=name,
                        **{
                            "Вид заготовки": kind or "Лист",
                            "Заготовка": blank,
                            "Размер заготовки": size,
                            "Материал": material,
                            "Гост материала": material_gost,
                            "Гост сортамента": sort_gost,
                            "Количество (норма расхода материала)": qty,
                            "Ед. измерения": unit,
                            "УТ код": source.get("УТ"),
                            "Номенклатура НСИ": source.get("НАИМЕНОВАНИЕ ЗАГОТОВКИ"),
                            "Комментарий": sheet_name,
                        },
                    )
                )
                continue
            elif "квадрат" in sheet_name.lower() or "листа" in sheet_name.lower():
                length = source.get("Требуемая длина заготовки, мм")
                width = source.get("Требуемая ширина заготовки, мм")
                height = source.get("Требуемая толщина заготовки, мм")
                kind = "Квадрат" if decimal_text(width) == decimal_text(height) and sort_gost else "Лист"
                size = f"W{decimal_text(height)} H{decimal_text(width)} L{decimal_text(length)}"
                blank = kind
                qty = source.get("Кол-во деталей, шт") or 1
                unit = "шт"
            else:
                kind = "Круг"
                diameter = source.get("Требуемый диаметр заготовки (Ø, мм)")
                size = f"D{decimal_text(diameter)}" if decimal_text(diameter) != "0" else ""
                blank = f"Круг {decimal_text(diameter)}" if size else "Круг"
                qty = consumption_from_length(source.get("Требуемая длина прутка (L,мм)") or source.get("Требуемая длина заготовки (L, мм)"))
                unit = "пог. м"

            rows.append(
                standard_row(
                    path.name,
                    IPS=ips,
                    Обозначение=source.get("Обозначение детали"),
                    Наименование=source.get("Наименование детали"),
                    **{
                        "Вид заготовки": kind,
                        "Заготовка": blank,
                        "Размер заготовки": size,
                        "Материал": material,
                        "Гост материала": material_gost,
                        "Гост сортамента": sort_gost,
                        "Количество (норма расхода материала)": qty,
                        "Ед. измерения": unit,
                        "УТ код": "",
                        "Номенклатура НСИ": "",
                        "Комментарий": sheet_name,
                    },
                )
            )
    return rows


def completeness_score(row: pd.Series) -> int:
    important = ["УТ код", "Заготовка", "Размер заготовки", "Материал", "Гост материала", "Гост сортамента", "Количество (норма расхода материала)"]
    score = sum(1 for column in important if text(row.get(column)) not in {"", "-", "0"})
    source = text(row.get("Источник файла"))
    if source == "Лист Microsoft Excel.xlsx":
        score += 30
    elif source == "Сопоставление_заказа_заготовок (1).xlsx":
        score += 20
    else:
        score += 10
    return score


def build_result() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    all_rows = []
    for file_path in INPUT_FILES:
        if file_path.name == "Лист Microsoft Excel.xlsx":
            all_rows.extend(load_prepared_library(file_path))
        elif file_path.name.startswith("Сопоставление_заказа_заготовок"):
            all_rows.extend(load_matching_result(file_path))
        else:
            all_rows.extend(load_raw_manufacturing(file_path))

    full = pd.DataFrame(all_rows, columns=TARGET_COLUMNS)
    full["IPS"] = full["IPS"].astype(str).str.strip()
    full = full[full["IPS"].ne("")]
    full["_score"] = full.apply(completeness_score, axis=1)
    full = full.sort_values(["IPS", "_score"], ascending=[True, False])
    duplicates = full[full.duplicated("IPS", keep="first")].copy()
    result = full.drop_duplicates("IPS", keep="first").drop(columns=["_score"]).sort_values("IPS")
    duplicates = duplicates.drop(columns=["_score"]).sort_values("IPS")

    output_file = OUTPUT_FILE
    try:
        result.to_excel(output_file, index=False, sheet_name="Итог")
    except PermissionError:
        output_file = OUTPUT_DIR / f"{OUTPUT_FILE.stem}_{datetime.now():%H%M%S}{OUTPUT_FILE.suffix}"
        result.to_excel(output_file, index=False, sheet_name="Итог")

    with pd.ExcelWriter(output_file, mode="a", engine="openpyxl", if_sheet_exists="replace") as writer:
        duplicates.to_excel(writer, index=False, sheet_name="Удаленные дубли")

    workbook = load_workbook(output_file)
    source_links = {path.name: str(path.resolve()) for path in INPUT_FILES}
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths = {
            "A": 38,
            "B": 13,
            "C": 28,
            "D": 34,
            "E": 18,
            "F": 30,
            "G": 22,
            "H": 24,
            "I": 20,
            "J": 20,
            "K": 18,
            "L": 14,
            "M": 18,
            "N": 48,
            "O": 28,
        }
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        for cell in worksheet["A"][1:]:
            target = source_links.get(text(cell.value))
            if target:
                cell.hyperlink = target
                cell.style = "Hyperlink"
    workbook.save(output_file)
    workbook.close()

    sample_path = OUTPUT_DIR / "sample.md"
    sample_path.write_text(to_markdown(result.head(7), TARGET_COLUMNS), encoding="utf-8")
    summary = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "source_files": [path.name for path in INPUT_FILES],
        "raw_rows": len(full),
        "result_rows": len(result),
        "duplicates_removed": len(duplicates),
        "output": str(output_file),
    }
    (OUTPUT_DIR / "summary.json").write_text(pd.Series(summary).to_json(force_ascii=False, indent=2), encoding="utf-8")
    return output_file


def to_markdown(frame: pd.DataFrame, columns: list[str]) -> str:
    rows = [[text(value) for value in frame.loc[:, columns].iloc[index].tolist()] for index in range(len(frame))]
    widths = [len(column) for column in columns]
    for row in rows:
        widths = [max(widths[index], len(row[index])) for index in range(len(columns))]
    header = "| " + " | ".join(column.ljust(widths[index]) for index, column in enumerate(columns)) + " |"
    sep = "| " + " | ".join("-" * widths[index] for index in range(len(columns))) + " |"
    body = ["| " + " | ".join(row[index].ljust(widths[index]) for index in range(len(columns))) + " |" for row in rows]
    return "\n".join([header, sep, *body])


if __name__ == "__main__":
    print(build_result())
