from __future__ import annotations

import csv
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(r"X:\19_МЕХ УЧАСТОК\База МСК")
CURRENT_MSK_DIR = BASE / "СПИСОК МСК"
OUTPUT_DIR = Path("reports/msk_excel_analysis_20260723")
APP_DB = Path("Данные для работы") / "BlankDemandPlanner.db"

KNOWN_BLANK_TYPES = [
    "Круг",
    "Лист",
    "Плита",
    "Поковка",
    "Труба",
    "Квадрат",
    "Шестигранник",
    "Уголок",
    "Швеллер",
    "Двутавр",
    "Винт",
    "Болт",
    "Гайка",
    "Шайба",
    "Втулка",
    "Прокат",
    "Профиль",
    "Заготовка",
    "Отливка",
    "Литье",
    "Пруток",
    "Стержень",
]

LONG_UNITS = {"Круг", "Труба", "Квадрат", "Шестигранник", "Уголок", "Швеллер", "Двутавр", "Пруток", "Стержень"}
PIECE_UNITS = {"Лист", "Плита", "Поковка", "Отливка", "Литье"}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def decimal_text(value: float) -> str:
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def parse_decimal(value: str) -> float:
    return float(value.replace(",", "."))


def gosts_from(text: str) -> list[str]:
    return re.findall(r"ГОСТ\s+[РРA-ZА-Яа-я0-9.\-\/ ]*?\d{2,5}(?:-\d{2,4})?", text, flags=re.IGNORECASE)


def first_gost(text: str) -> str:
    found = gosts_from(text)
    return clean(found[0]) if found else ""


def strip_gosts(text: str) -> str:
    result = re.sub(r"ГОСТ\s+[РРA-ZА-Яа-я0-9.\-\/ ]*?\d{2,5}(?:-\d{2,4})?", "", text, flags=re.IGNORECASE)
    return clean(result)


def split_material(material_text: str) -> tuple[str, str]:
    material_gost = first_gost(material_text)
    material = strip_gosts(material_text)
    return material, material_gost


def infer_blank_type(blank_text: str) -> str:
    normalized = blank_text.strip()
    for blank_type in KNOWN_BLANK_TYPES:
        if re.search(rf"(^|\s){re.escape(blank_type)}(\s|$)", normalized, re.IGNORECASE):
            return blank_type
    return normalized.split()[0] if normalized else ""


def looks_like_size(text: str) -> bool:
    return bool(re.search(r"(?:[ØФфDd]\s*\d|\d+\s*[xх×]\s*\d|L\s*=|\bмм\b)", text, re.IGNORECASE))


def split_blank_size_and_profile_gost(blank_text: str, size_text: str, blank_type: str) -> tuple[str, str, str]:
    blank = clean(blank_text)
    size = clean(size_text)

    parts = [part.strip() for part in re.split(r"\s*;\s*", blank) if part.strip()]
    if len(parts) > 1 and looks_like_size(parts[-1]):
        size = size or parts[-1]
        blank = "; ".join(parts[:-1])

    profile_gost = first_gost(blank)
    blank = strip_gosts(blank)
    blank = remove_profile_class(blank)
    return blank, normalize_size_for_library(size, blank_type), profile_gost


def remove_profile_class(text: str) -> str:
    result = clean(text)
    result = re.sub(r"(\d+(?:[,.]\d+)?)\s*[-–]\s*[А-ЯA-Z]\d?\b", r"\1", result, flags=re.IGNORECASE)
    result = re.sub(r"\b[А-ЯA-Z]\d?\b$", "", result, flags=re.IGNORECASE)
    return clean(result)


def normalize_size_full(size_text: str, blank_type: str) -> str:
    text = clean(size_text).replace("×", "х").replace("X", "х").replace("x", "х")
    if not text:
        return ""

    nums = [parse_decimal(item) for item in re.findall(r"\d+(?:[,.]\d+)?", text)]
    diameter = re.search(r"[ØФфDd]\s*=?\s*(\d+(?:[,.]\d+)?)", text)
    length = re.search(r"\bL\s*=?\s*(\d+(?:[,.]\d+)?)", text, re.IGNORECASE)

    if blank_type in LONG_UNITS and (diameter or length):
        parts = []
        if diameter:
            parts.append(f"D{decimal_text(parse_decimal(diameter.group(1)))}")
        if length:
            parts.append(f"L{decimal_text(parse_decimal(length.group(1)))}")
        return " ".join(parts) if parts else text

    if (blank_type in {"Лист", "Плита"} or re.search(r"\d+\s*х\s*\d+\s*х\s*\d+", text)) and len(nums) >= 3:
        length_value, width_value, height_value = nums[0], nums[1], nums[2]
        return f"W{decimal_text(height_value)} H{decimal_text(width_value)} L{decimal_text(length_value)}"

    return text


def normalize_size_for_library(size_text: str, blank_type: str) -> str:
    full = normalize_size_full(size_text, blank_type)
    if blank_type in LONG_UNITS:
        diameter = re.search(r"\bD\d+(?:[,.]\d+)?", full, re.IGNORECASE)
        return diameter.group(0).replace(",", ".") if diameter else full
    return full


def infer_consumption_and_unit(blank_type: str, raw_size_text: str) -> tuple[str, str]:
    full_size = normalize_size_full(raw_size_text, blank_type)
    length = re.search(r"\bL(\d+(?:[,.]\d+)?)", full_size, re.IGNORECASE)
    if blank_type in LONG_UNITS and length:
        mm = parse_decimal(length.group(1))
        return decimal_text(mm / 1000), "пог. м"
    if blank_type in PIECE_UNITS:
        return "1", "шт"
    if blank_type and blank_type not in LONG_UNITS:
        return "1", "шт"
    return "", ""


def is_current_msk(path: Path) -> bool:
    try:
        path.relative_to(CURRENT_MSK_DIR)
        return True
    except ValueError:
        return False


def extract_msk(path: Path) -> tuple[dict[str, Any] | None, str | None]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return None, f"{type(exc).__name__}: {exc}"

    try:
        ws = wb[wb.sheetnames[0]]
        form = clean(ws["A1"].value)
        if not form.startswith("Форма"):
            return None, "not_msk_form"

        raw_blank = clean(ws["D7"].value)
        raw_size = clean(ws["D8"].value)
        raw_material = clean(ws["D6"].value)
        blank_type = infer_blank_type(raw_blank)
        blank, normalized_size, profile_gost = split_blank_size_and_profile_gost(raw_blank, raw_size, blank_type)
        material, material_gost = split_material(raw_material)
        consumption, unit = infer_consumption_and_unit(blank_type, raw_size)
        source_ut_code = clean(ws["O7"].value) if clean(ws["J7"].value).lower().startswith("код ут") else ""

        return {
            "Источник": str(path),
            "_is_current": is_current_msk(path),
            "_mtime": path.stat().st_mtime,
            "IPS": clean(ws["C5"].value),
            "Обозначение": clean(ws["E5"].value),
            "Наименование": clean(ws["L5"].value),
            "Вид заготовки": blank_type,
            "Заготовка": blank,
            "Размер заготовки": normalized_size,
            "Материал": material,
            "Гост материала": material_gost,
            "Гост сортамента": profile_gost,
            "Количество (норма расхода материала)": consumption,
            "Ед. измерения": unit,
            "Вариант УТ кода": source_ut_code,
            "Номенклатура НСИ": "",
            "Разработчик МСК": clean(ws["X3"].value),
            "Дата МСК": clean(ws["AA3"].value),
            "Проблемы качества": "",
        }, None
    finally:
        wb.close()


def quality_issues(row: dict[str, Any]) -> str:
    issues = []
    if not row["IPS"] or row["IPS"].lower() == "нет ипс":
        issues.append("нет IPS")
    if not row["Обозначение"]:
        issues.append("нет обозначения")
    if not row["Наименование"]:
        issues.append("нет наименования")
    if not row["Заготовка"]:
        issues.append("нет заготовки")
    if not row["Материал"]:
        issues.append("нет материала")
    if not row["Гост материала"]:
        issues.append("нет ГОСТ материала")
    if not row["Гост сортамента"]:
        issues.append("нет ГОСТ сортамента")
    if not row["Размер заготовки"]:
        issues.append("нет размера заготовки")
    if not row["Количество (норма расхода материала)"]:
        issues.append("нет расчетной нормы расхода")
    if not row["Ед. измерения"]:
        issues.append("нет единицы измерения")
    if not row["Вариант УТ кода"]:
        issues.append("нет варианта УТ")
    return "; ".join(issues)


def dedupe_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    passthrough: list[dict[str, Any]] = []
    for row in rows:
        ips = str(row["IPS"]).strip()
        if ips and ips.lower() != "нет ипс":
            groups[ips].append(row)
        else:
            passthrough.append(row)

    selected: list[dict[str, Any]] = []
    duplicate_rows: list[dict[str, Any]] = []
    for ips, group in groups.items():
        ranked = sorted(group, key=lambda r: (bool(r["_is_current"]), float(r["_mtime"])), reverse=True)
        chosen = ranked[0]
        selected.append(chosen)
        for other in ranked[1:]:
            duplicate_rows.append(
                {
                    "IPS": ips,
                    "Выбран источник": chosen["Источник"],
                    "Отклонен источник": other["Источник"],
                    "Причина": "выбран актуальный СПИСОК МСК" if chosen["_is_current"] and not other["_is_current"] else "выбран более свежий файл",
                }
            )

    selected.extend(passthrough)
    selected.sort(key=lambda r: (not bool(r["_is_current"]), str(r["IPS"]), str(r["Обозначение"])))
    return selected, duplicate_rows


def normalize_match_text(value: str) -> str:
    return re.sub(r"[^0-9a-zа-я]+", " ", value.lower()).strip()


def normalize_material_key(value: str) -> str:
    material, _ = split_material(value)
    return normalize_match_text(material)


def size_tokens(size_text: str) -> set[str]:
    return set(re.findall(r"\b[DWHTSL]\d+(?:[,.]\d+)?", size_text.upper()))


def load_ut_candidates() -> list[dict[str, str]]:
    if not APP_DB.exists():
        return []
    query = """
        select
            ba.OneCCode,
            ba.SourceName,
            cb.Material,
            cb.DiameterMm,
            cb.WidthMm,
            cb.HeightMm,
            cb.ThicknessMm,
            cb.WallThicknessMm,
            cb.LengthMm
        from BlankAliases ba
        left join CanonicalBlanks cb on cb.Id = ba.CanonicalBlankId
        where ba.OneCCode is not null and trim(ba.OneCCode) <> ''
          and ba.IsActive = 1
    """
    try:
        with sqlite3.connect(APP_DB) as connection:
            connection.row_factory = sqlite3.Row
            candidates = []
            for row in connection.execute(query):
                size = " ".join(
                    part
                    for part in [
                        f"D{decimal_text(float(row['DiameterMm']))}" if clean(row["DiameterMm"]) else "",
                        f"W{decimal_text(float(row['WidthMm']))}" if clean(row["WidthMm"]) else "",
                        f"H{decimal_text(float(row['HeightMm']))}" if clean(row["HeightMm"]) else "",
                        f"T{decimal_text(float(row['ThicknessMm']))}" if clean(row["ThicknessMm"]) else "",
                        f"S{decimal_text(float(row['WallThicknessMm']))}" if clean(row["WallThicknessMm"]) else "",
                        f"L{decimal_text(float(row['LengthMm']))}" if clean(row["LengthMm"]) else "",
                    ]
                    if part
                )
                source_name = clean(row["SourceName"])
                candidates.append(
                    {
                        "Код УТ": clean(row["OneCCode"]),
                        "Номенклатура НСИ": source_name,
                        "Материал НСИ": clean(row["Material"]),
                        "Размер НСИ": size,
                        "Вид НСИ": infer_blank_type(source_name),
                    }
                )
            return candidates
    except (sqlite3.Error, ValueError):
        return []


def best_ut_candidate(row: dict[str, Any], candidates: list[dict[str, str]]) -> tuple[str, str]:
    if row["Вариант УТ кода"]:
        return row["Вариант УТ кода"], row["Номенклатура НСИ"]

    row_type = row["Вид заготовки"]
    row_material = normalize_material_key(row["Материал"])
    row_sizes = size_tokens(row["Размер заготовки"])
    if not row_type or not row_material or not row_sizes:
        return "", ""

    scored = []
    for candidate in candidates:
        if candidate["Вид НСИ"] != row_type:
            continue
        if normalize_material_key(candidate["Материал НСИ"]) != row_material:
            continue
        candidate_sizes = size_tokens(candidate["Размер НСИ"])
        if not row_sizes.issubset(candidate_sizes):
            continue
        score = len(row_sizes) * 20
        score += 10 if row["Гост сортамента"] and row["Гост сортамента"] in candidate["Номенклатура НСИ"] else 0
        score += 10 if row["Гост материала"] and row["Гост материала"] in candidate["Номенклатура НСИ"] else 0
        scored.append((score, candidate))

    if not scored:
        return "", ""

    _, candidate = sorted(scored, key=lambda item: (item[0], item[1]["Код УТ"]), reverse=True)[0]
    return candidate["Код УТ"], candidate["Номенклатура НСИ"]


def fill_ut_suggestions(rows: list[dict[str, Any]], candidates: list[dict[str, str]]) -> int:
    count = 0
    for row in rows:
        code, name = best_ut_candidate(row, candidates)
        if code:
            row["Вариант УТ кода"] = code
            row["Номенклатура НСИ"] = name
            count += 1
        row["Проблемы качества"] = quality_issues(row)
    return count


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_sheet(ws: Any, rows: list[dict[str, Any]], columns: list[str], hyperlink_source: bool = False) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
        if hyperlink_source and "Источник" in columns:
            source_col = columns.index("Источник") + 1
            cell = ws.cell(ws.max_row, source_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = [
        path
        for path in BASE.rglob("*")
        if path.is_file()
        and path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
        and not path.name.startswith("~$")
    ]

    raw_rows: list[dict[str, Any]] = []
    skipped = Counter()
    for path in files:
        row, error = extract_msk(path)
        if row:
            raw_rows.append(row)
        elif error:
            skipped[error] += 1

    rows, duplicates = dedupe_rows(raw_rows)
    candidates = load_ut_candidates()
    ut_suggestions = fill_ut_suggestions(rows, candidates)

    main_columns = [
        "Источник",
        "IPS",
        "Обозначение",
        "Наименование",
        "Вид заготовки",
        "Заготовка",
        "Размер заготовки",
        "Материал",
        "Гост материала",
        "Гост сортамента",
        "Количество (норма расхода материала)",
        "Ед. измерения",
        "Вариант УТ кода",
        "Номенклатура НСИ",
        "Разработчик МСК",
        "Дата МСК",
        "Проблемы качества",
    ]
    duplicate_columns = ["IPS", "Выбран источник", "Отклонен источник", "Причина"]

    csv_path = OUTPUT_DIR / "msk_library_extract.csv"
    write_csv(csv_path, rows, main_columns)

    old_suggestions_csv = OUTPUT_DIR / "msk_ut_code_suggestions.csv"
    if old_suggestions_csv.exists():
        old_suggestions_csv.unlink()

    xlsx_path = OUTPUT_DIR / "msk_library_extract.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Библиотека МСК"
    write_sheet(ws, rows, main_columns, hyperlink_source=True)
    widths = {
        "A": 70,
        "B": 13,
        "C": 28,
        "D": 34,
        "E": 18,
        "F": 32,
        "G": 22,
        "H": 24,
        "I": 20,
        "J": 20,
        "K": 18,
        "L": 14,
        "M": 18,
        "N": 55,
        "O": 20,
        "P": 14,
        "Q": 46,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    duplicate_ws = wb.create_sheet("Дубли IPS")
    write_sheet(duplicate_ws, duplicates, duplicate_columns)
    for letter, width in {"A": 13, "B": 70, "C": 70, "D": 30}.items():
        duplicate_ws.column_dimensions[letter].width = width

    try:
        wb.save(xlsx_path)
    except PermissionError:
        xlsx_path = OUTPUT_DIR / f"msk_library_extract_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        wb.save(xlsx_path)
    wb.close()

    quality = Counter()
    for row in rows:
        if row["Проблемы качества"]:
            for issue in row["Проблемы качества"].split("; "):
                quality[issue] += 1

    summary = {
        "base": str(BASE),
        "files_seen": len(files),
        "msk_rows_extracted_raw": len(raw_rows),
        "msk_rows_after_ips_dedupe": len(rows),
        "duplicates_removed": len(duplicates),
        "skipped": dict(skipped),
        "current_spisok_rows_after_dedupe": sum(1 for row in rows if row["_is_current"]),
        "blank_types_top": dict(Counter(row["Вид заготовки"] or "(пусто)" for row in rows).most_common(30)),
        "quality_issues": dict(quality),
        "ut_suggestions": ut_suggestions,
        "csv": str(csv_path),
        "xlsx": str(xlsx_path),
    }

    summary_path = OUTPUT_DIR / "summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
